"""販売図面から読み取った物件情報（JSON）をもとに、資金計算書（概算）の Excel を作る。

使い方:
    python3 shohiyo/make_estimate.py 物件.json [出力.xlsx]

JSON の項目は README.md を参照。概算値（登記費用・火災保険など）は ESTIMATES の
ルールで自動設定し、JSON の "overrides" で個別に上書きできる。
"""
import datetime
import json
import re
import sys
import warnings
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, Side

sys.path.insert(0, str(Path(__file__).resolve().parent))
from reds_fee import reds_fee  # noqa: E402

TYPES = ('新築戸建', '中古戸建', '新築マンション', '中古マンション', '土地')
DEALS = ('仲介', '売主', '代理')
FEE_PLANS = ('割引', '無料', '半額')
STAFF = '柴田'
LOAN_RATE = 1.195  # 標準の金利（%）
OTHER_AGENT_ADMIN_FEE = 55000  # 他社の事務代行手数料（税込）
BANK_FEE_RATE = 0.022  # 融資事務手数料（融資額 × 2.2%）

FONT = '游ゴシック'
YEN = '"¥"#,##0;[Red]"¥"\\-#,##0'
RED = 'FFFF0000'
BLUE = 'FF0000FF'
thin = Side(style='thin', color='FF999999')


# ---- 概算ルール（図面からは分からない費用） -------------------------------------

THIS_YEAR = datetime.date.today().year


def building_age(s):
    if s['type'].startswith('新築'):
        return 0
    return THIS_YEAR - s['built_year'] if s.get('built_year') else 20


def is_wood(s):
    return '木' in (s.get('structure') or ('木造' if s['type'].endswith('戸建') else 'RC'))


def prefecture(address):
    m = re.match(r'(東京都|北海道|(?:京都|大阪)府|.{2,3}県)', address or '')
    return m.group(1) if m else ''


# ---- 登記費用（固定資産税評価額から計算） ---------------------------------------
# 評価額が図面に無い場合の推定。分かれば JSON の assessed_land / assessed_building で指定する。

def estimate_assessed_land(s):
    """土地（マンションは敷地権の持分）の固定資産税評価額の推定。"""
    ratio = {'新築マンション': 0.08, '中古マンション': 0.08, '新築戸建': 0.35, '中古戸建': 0.40, '土地': 0.70}
    return round(s['price'] * ratio[s['type']], -4)


def estimate_assessed_building(s):
    """建物の固定資産税評価額の推定（新築は法務局の認定価格相当）。床面積 × 再建築費単価 × 経年残価率。"""
    if s['type'] == '土地':
        return 0
    age = building_age(s)
    if s['type'].endswith('マンション'):
        area, unit, decay = s.get('floor_area') or 70, 150000, 0.015
    elif is_wood(s):
        area, unit, decay = s.get('floor_area') or 100, 110000, 0.045
    else:
        area, unit, decay = s.get('floor_area') or 100, 150000, 0.015
    return round(area * unit * max(0.2, 1 - decay * age), -4)


def registration_rows(s):
    """登記費用の内訳: (登記の種類, 課税標準(値 or 式), 税率, 根拠)。"""
    rows = [('土地 所有権移転', s.get('assessed_land', estimate_assessed_land(s)), 0.015,
             '土地の売買による所有権移転登記の軽減税率1.5%（本則2%）')]
    if s['type'] != '土地':
        area_ok = (s.get('floor_area') or 50) >= 50
        new_seismic = s['type'].startswith('新築') or not is_old_seismic(s.get('built_year'))
        if s['type'].startswith('新築'):
            rows.append(('建物 所有権保存', s.get('assessed_building', estimate_assessed_building(s)),
                         0.0015 if area_ok else 0.004, '住宅用家屋の所有権保存登記 0.15%（本則0.4%）'))
        else:
            rate = 0.003 if area_ok and new_seismic else 0.02
            rows.append(('建物 所有権移転', s.get('assessed_building', estimate_assessed_building(s)), rate,
                         '住宅用家屋の所有権移転登記 0.3%（床面積50㎡以上・新耐震。該当しなければ本則2%）'))
    rows.append(('抵当権設定', '=E@LOAN@', 0.001, '住宅用家屋の抵当権設定登記 0.1%（本則0.4%）。課税標準は融資額'))
    return rows


JUDICIAL_SCRIVENER_FEE = 120000  # 司法書士報酬・登記事項証明等の実費（税込の目安）


# ---- 火災保険（5年・地震保険込の相場） ------------------------------------------
# 年額保険料（保険金額1,000万円あたり）の目安
FIRE_RATE = {'M': 5000, 'T': 9000, 'H': 16000}
QUAKE_RATE = {  # 地震保険（イ構造, ロ構造）。2022年10月改定の基本料率の目安
    '東京都': (27500, 42200), '神奈川県': (27500, 42200), '千葉県': (27500, 42200), '静岡県': (27500, 42200),
}
QUAKE_RATE_DEFAULT = (11600, 23000)
LONG_TERM_FIRE, LONG_TERM_QUAKE = 4.7, 4.65  # 5年一括の係数
QUAKE_DISCOUNT = 0.9  # 建築年割引（1981年6月以降の建物）10%


def estimate_fire_insurance(s):
    """(保険料, 根拠メモ)。建物のみ（家財なし）、地震保険金額は火災の50%。高めに見て10万円単位で繰り上げ。"""
    kind = s['type']
    if kind == '土地':
        return 0, ''
    if kind.endswith('マンション'):
        area, unit, cls = s.get('floor_area') or 70, 200000, 'M'
    elif is_wood(s):
        area, unit = s.get('floor_area') or 100, 180000
        cls = 'T' if kind.startswith('新築') else 'H'
    else:
        area, unit, cls = s.get('floor_area') or 100, 220000, 'T'
    insured = int(round(area * unit, -5))
    quake_cls = 0 if cls in ('M', 'T') else 1
    pref = prefecture(s.get('address'))
    quake_rate = QUAKE_RATE.get(pref, QUAKE_RATE_DEFAULT)[quake_cls]
    fire = insured / 1e7 * FIRE_RATE[cls] * LONG_TERM_FIRE
    discount = QUAKE_DISCOUNT if not is_old_seismic(s.get('built_year')) else 1
    quake = insured * 0.5 / 1e7 * quake_rate * discount * LONG_TERM_QUAKE
    total = -(-int(fire + quake) // 100000) * 100000
    note = (f"相場目安: {cls}構造・建物保険金額{insured // 10000:,}万円（{area}㎡）、5年一括。"
            f"火災 約{int(fire):,}円＋地震（{pref or '所在地不明'}・{'イロ'[quake_cls]}構造・50%）約{int(quake):,}円。"
            "家財は含まない。10万円単位で繰り上げ")
    return total, note


def estimate_property_tax(kind):
    """固定資産税・都市計画税の日割清算金の概算。"""
    return {'土地': 50000, '新築マンション': 50000, '中古マンション': 80000}.get(kind, 100000)


def is_old_seismic(built_year):
    """旧耐震（1981年5月以前の建築確認）の可能性があるか。築年のみで判定する。"""
    return built_year is not None and built_year <= 1981


# ---- 物件情報の検証・既定値 ---------------------------------------------------

def normalize(spec):
    s = dict(spec)
    s['_given'] = tuple(spec)
    if s.get('type') not in TYPES:
        raise ValueError(f"type は {TYPES} のいずれか: {s.get('type')!r}")
    if s.get('deal', '仲介') not in DEALS:
        raise ValueError(f"deal は {DEALS} のいずれか: {s.get('deal')!r}")
    s['deal'] = s.get('deal', '仲介')
    s['price'] = int(s['price'])
    # 仲介手数料の既定: 図面が「売主」「代理」、または「手数料3%」等（売主側が手数料を負担）の記載なら無料。
    # 「仲介」「媒介」「専任」「専属」なら割引。
    if 'fee_plan' not in s:
        s['fee_plan'] = '無料' if s['deal'] in ('売主', '代理') or s.get('fee_3pct') else '割引'
    s.setdefault('staff', STAFF)
    if s['fee_plan'] not in FEE_PLANS:
        raise ValueError(f"fee_plan は {FEE_PLANS} のいずれか: {s['fee_plan']!r}")
    loan = {'amount': None, 'rate': LOAN_RATE, 'years': 35, 'bonus': 0,
            'bank': '（仮）都市銀行', 'rate_type': '変動金利'}
    loan.update(s.get('loan') or {})
    if loan['amount'] is None:
        loan['amount'] = s['price']  # 既定: 物件価格を全額借入、諸費用は自己資金
    s['loan'] = loan
    s.setdefault('overrides', {})
    return s


def build_items(s):
    """諸経費の行リスト。各行: (項目名, 説明, REDS側の値/式, 他社側の値/式, 約をつけるか, 根拠コメント)"""
    price, kind, ov = s['price'], s['type'], s['overrides']
    items = []

    def val(key, default):
        return ov.get(key, default)

    items.append(('契約書印紙代',
                  '（売買契約書の印紙税。軽減税率：1千万超5千万以下1万円、5千万超1億以下3万円、1億超5億以下6万円）',
                  '=LOOKUP(E8-1,{0,500000,1000000,5000000,10000000,50000000,100000000,500000000},'
                  '{200,500,1000,5000,10000,30000,60000,160000})',
                  '=E@ROW@', False, '印紙税法 別表第一 第1号文書・租税特別措置法の軽減税率（2027年3月31日まで）'))

    legal = '=INT(IF(E8<=2000000,INT(E8*0.05),IF(E8<=4000000,INT(100000+(E8-2000000)*0.04),INT(E8*0.03+60000)))*1.1)'
    plan = s['fee_plan']
    if plan == '無料':
        reds = 0
        reason = '取引態様が' + s['deal'] if s['deal'] != '仲介' else '図面に手数料3%等の記載'
        note = f"{reason}のため無料" if 'fee_plan' not in s.get('_given', ()) else None
    elif plan == '半額':
        reds = '=INT(H@ROW@/2)'
        note = '他社（法定上限）の半額'
    else:
        r = reds_fee(price)
        reds = r['fee_and_tax']
        note = (f"REDS手数料計算（calc10_2）で算出: ({r['rate']}%＋3万円)×1.1"
                + (f"、上限 {r['limit_str']}" if r['limit_str'] else '')
                + f"／割引率 {r['discount_rate']}%")
    items.append((f'仲介手数料（{plan}）', '（一般的な仲介手数料は、物件価格の3％＋6万円に消費税）',
                  val('仲介手数料', reds), legal, False, note))
    items.append(('事務代行手数料', '（契約書類作成、物件調査、住宅ローン等の代行手数料　約5～10万円）',
                  0, OTHER_AGENT_ADMIN_FEE, True, None))

    items.append(('登記費用', '（移転登記・保存登記・抵当権設定の登録免許税と司法書士報酬。内訳は下記）',
                  val('登記費用', '=ROUNDUP(E@REG@,-5)'), '=E@ROW@', True,
                  '下記「登記費用の内訳」の合計を10万円単位で繰り上げ'))
    if kind != '土地':
        premium, note = estimate_fire_insurance(s)
        items.append(('火災保険', '（火災保険5年加入・地震保険込の相場）',
                      val('火災保険', premium), '=E@ROW@', True, note))
    items.append(('固定資産税・都市計画税清算金', '（固定資産税・都市計画税の年間支払額を引渡日で日割清算します）',
                  val('固定資産税・都市計画税清算金', estimate_property_tax(kind)), '=E@ROW@', True,
                  '概算（目安）。年税額が分かれば日割で差し替え'))
    if kind == '新築戸建':
        items.append(('表示登記費用', '（建物表題登記。土地家屋調査士に支払います）',
                      val('表示登記費用', 150000), '=E@ROW@', True, '概算（目安）'))
    if kind.endswith('マンション'):
        mf, rr = s.get('management_fee') or 0, s.get('repair_reserve') or 0
        settle = f'=ROUNDUP(({mf}+{rr})*2,-3)' if mf or rr else 60000
        items.append(('管理費・修繕積立金清算金', '（管理費・修繕積立金の日割清算。約2ヶ月分）',
                      val('管理費・修繕積立金清算金', settle), '=E@ROW@', True,
                      f"図面の管理費{mf:,}円＋修繕積立金{rr:,}円 × 2ヶ月（百の位を繰り上げ）" if mf or rr
                      else '概算（図面に記載なし）'))
    if kind == '新築マンション' and s.get('initial_repair_fund'):
        items.append(('修繕積立基金', '（新築時に一括で支払う修繕積立基金）',
                      s['initial_repair_fund'], '=E@ROW@', False, '販売図面記載額'))
    items.append(('融資事務手数料', '（銀行に支払う融資手数料です。融資額×2.2％）',
                  f'=INT(E@LOAN@*{BANK_FEE_RATE})', '=E@ROW@', True, '融資事務手数料型（定率2.2%）の場合'))
    items.append(('金銭消費貸借契約書印紙代', '（銀行との金銭消費貸借契約書に貼付する収入印紙。WEB契約の場合は無料）',
                  '=IF(E@LOAN@<=0,0,LOOKUP(E@LOAN@-1,{0,100000,500000,1000000,5000000,10000000,50000000,100000000},'
                  '{200,400,1000,2000,10000,20000,60000,100000}))',
                  '=E@ROW@', False, '印紙税法 別表第一 第1号文書（消費貸借契約書）'))
    return items


# ---- Excel 出力 --------------------------------------------------------------

def write_workbook(s, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = '資金計算書'
    for col, w in {'A': 14, 'B': 20, 'C': 23, 'D': 10, 'E': 24, 'F': 3, 'G': 5, 'H': 22}.items():
        ws.column_dimensions[col].width = w

    def put(ref, value, size=12, bold=False, color=None, fmt=None, align=None, wrap=False):
        c = ws[ref]
        c.value = value
        c.font = Font(name=FONT, size=size, bold=bold, color=color)
        if fmt:
            c.number_format = fmt
        if align or wrap:
            c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
        return c

    # 見出し
    ws.merge_cells('C1:E2')
    put('C1', '資金計算書（概算）', 16, True, align='center')
    ws.merge_cells('F1:H1')
    put('F1', '=TODAY()', 10, True, fmt='[$-411]ggge"年"m"月"d"日";@', align='center')
    ws.merge_cells('F2:H2')
    put('F2', '株式会社不動産流通システム', 10, True, align='center')
    ws.merge_cells('F3:H3')
    put('F3', f"担当：{s.get('staff', '')}", 10, True, align='center')

    put('A4', '物件名', bold=True)
    ws.merge_cells('B4:H4')
    put('B4', s.get('name', ''), align='left')
    put('A5', '物件所在地', bold=True)
    ws.merge_cells('B5:H5')
    put('B5', s.get('address', ''), align='left')

    put('A8', '物件価格', bold=True)
    put('C8', '売買代金', 12, align='right')
    put('E8', s['price'], 14, True, BLUE, YEN, 'right')
    put('C10', '小計', 12, align='right')
    put('E10', '=E8', 14, True, fmt=YEN, align='right')

    put('A13', '諸経費', bold=True)
    put('E13', 'REDSの場合', bold=True, align='center')
    put('H13', '他社の場合', bold=True, align='center')

    items = build_items(s)
    first = row = 14
    # 融資金額のセル位置は諸経費の行数で決まるので先に計算する
    sub_row = first + len(items) * 2
    total_row, diff_row = sub_row + 2, sub_row + 4
    own_row, loan_row = sub_row + 6, sub_row + 8
    notes = ['※当該計算書の金額は作成時での概算のものであり、実際のお支払い額と異なる場合がございます。',
             '※上記費用の他、不動産取得後に不動産取得税の納税が必要となる場合があります。納期は各都道府県により異なります。']
    if is_old_seismic(s.get('built_year')):
        notes.append('※耐震基準適合証明書の発行（費用約5～6万円）を受けると、住宅ローン減税、不動産取得税・登録免許税の'
                     '減税の適用を受けることができる場合がございます。')
    notes_row = loan_row + 10  # 住宅ローン明細の下に注記
    print_end = notes_row + len(notes) - 1
    reg = registration_rows(s)
    reg_head = print_end + 3  # 登記費用の内訳は印刷範囲の外
    reg_total = reg_head + 1 + len(reg) + 1
    for name, desc, reds, other, approx, note in items:
        fill = lambda v: (v.replace('@ROW@', str(row)).replace('@LOAN@', str(loan_row))
                          .replace('@REG@', str(reg_total))) if isinstance(v, str) else v
        ws.merge_cells(f'B{row}:C{row}')
        put(f'B{row}', name, 14, True)
        if approx:
            put(f'D{row}', '約', 12, align='right')
            put(f'G{row}', '約', 12, align='right')
        is_fee = name.startswith('仲介手数料')
        reds_val, other_val = fill(reds), fill(other)
        e = put(f'E{row}', reds_val, 14, True, RED if is_fee else (BLUE if not isinstance(reds, str) else None),
                YEN, 'right')
        put(f'H{row}', other_val, 14, True, RED if is_fee else None, YEN, 'right')
        if note:
            e.comment = Comment(note, 'shohiyo')
        ws.merge_cells(f'B{row + 1}:H{row + 1}')
        put(f'B{row + 1}', desc, 10)
        row += 2

    put(f'C{sub_row}', '小計', 12, align='right')
    put(f'E{sub_row}', f'=SUM(E{first}:E{sub_row - 1})', 14, True, fmt=YEN, align='right')
    put(f'H{sub_row}', f'=SUM(H{first}:H{sub_row - 1})', 14, True, fmt=YEN, align='right')
    for col in 'BCDEFGH':
        ws[f'{col}{sub_row}'].border = Border(top=thin)

    put(f'A{total_row}', '総合計額', bold=True)
    put(f'C{total_row}', '総額', 12, align='right')
    put(f'E{total_row}', f'=E10+E{sub_row}', 14, True, fmt=YEN, align='right')
    put(f'H{total_row}', f'=E10+H{sub_row}', 14, True, fmt=YEN, align='right')
    put(f'E{diff_row}', '総額差額', 12, True, align='right')
    put(f'H{diff_row}', f'=H{total_row}-E{total_row}', 14, True, RED, YEN, 'right')

    loan = s['loan']
    put(f'A{own_row}', '自己資金', bold=True)
    put(f'C{own_row}', '（総額 − 融資金額）', 10)
    put(f'E{own_row}', f'=E{total_row}-E{loan_row}', 14, True, fmt=YEN, align='right')

    ws.merge_cells(f'A{loan_row}:B{loan_row}')
    put(f'A{loan_row}', '住宅ローン明細', bold=True)
    put(f'C{loan_row}', '融資金額', 12, align='right')
    c = put(f'E{loan_row}', loan['amount'], 14, True, BLUE, YEN, 'right')
    c.comment = Comment('既定は物件価格の全額借入（諸費用は自己資金）。変更すると自己資金・返済額が再計算されます', 'shohiyo')
    r = loan_row + 1
    ws.merge_cells(f'A{r}:H{r}')
    rate_row = loan_row + 4  # 下の表の金利・期間の行
    bank, rate_type = loan['bank'], loan['rate_type']
    put(f'A{r}', f'="{bank}（融資事務手数料型）：借入期間"&E{rate_row}&"年間、{rate_type}年"'
                 f'&TEXT(D{rate_row},"0.000")&"％　元利均等返済"', 11)
    ws.merge_cells(f'A{r + 1}:H{r + 1}')
    put(f'A{r + 1}', '※金利は金融機関・審査により異なります。一例です。', 10)

    h = r + 2
    for col, label in zip('BCDE', ('金融機関', '融資金額', '金利(%)', '期間(年)')):
        put(f'{col}{h}', label, 11, True, align='center')
    ws.merge_cells(f'F{h}:H{h}')
    put(f'F{h}', '月々均等のお支払い', 11, True, align='center')
    b = h + 1
    put(f'B{b}', loan['bank'].replace('（仮）', ''), 11, align='center')
    put(f'C{b}', f'=E{loan_row}', 12, True, fmt=YEN, align='right')
    put(f'D{b}', loan['rate'], 12, True, BLUE, '0.000', 'center')  # 四捨五入表示にしない
    put(f'E{b}', loan['years'], 12, True, BLUE, align='center')
    ws.merge_cells(f'F{b}:H{b}')
    put(f'F{b}', f'=IF(AND(C{b}>0,D{b}>0,E{b}>0),ROUNDUP(INT((C{b}*(D{b}/1200)*(1+D{b}/1200)^(E{b}*12))'
                 f'/((1+D{b}/1200)^(E{b}*12)-1)),0),"-")', 14, True, fmt=YEN, align='right')

    y = b + 2
    put(f'C{y}', '年間返済額', 12, align='right')
    put(f'D{y}', '12ヶ月', 10, align='center')
    put(f'E{y}', f'=IF(ISNUMBER(F{b}),F{b}*12,0)', 14, True, fmt=YEN, align='right')
    put(f'C{y + 1}', '月々のご返済', 12, align='right')
    put(f'E{y + 1}', f'=(E{y}-E{y + 2}*2)/12', 14, True, fmt=YEN, align='right')
    put(f'C{y + 2}', 'ボーナス加算', 12, align='right')
    put(f'D{y + 2}', '年2回', 10, align='center')
    put(f'E{y + 2}', loan['bonus'], 14, True, BLUE, YEN, 'right')

    assert notes_row == y + 4
    for i, text in enumerate(notes):
        n = notes_row + i
        ws.merge_cells(f'A{n}:H{n}')
        put(f'A{n}', text, 9, wrap=True)
        ws.row_dimensions[n].height = 26

    # 登記費用の内訳（固定資産税評価額から計算）。印刷範囲の外
    ws.merge_cells(f'A{reg_head}:B{reg_head}')
    put(f'A{reg_head}', '登記費用の内訳', bold=True)
    for col, label in zip('CDE', ('課税標準(評価額等)', '税率', '登録免許税')):
        put(f'{col}{reg_head}', label, 10, True, align='center')
    r = reg_head + 1
    for label, base, rate, why in reg:
        put(f'B{r}', label, 11)
        base_val = base.replace('@LOAN@', str(loan_row)) if isinstance(base, str) else base
        c = put(f'C{r}', base_val, 11, False, None if isinstance(base, str) else BLUE, YEN, 'right')
        if not isinstance(base, str):
            given = ('assessed_land' if label.startswith('土地') else 'assessed_building') in s
            c.comment = Comment('固定資産税評価額（評価証明書の値）' if given else
                                '固定資産税評価額の推定値です。評価証明書・公課証明の値が分かれば書き換えてください', 'shohiyo')
        d = put(f'D{r}', rate, 11, False, BLUE, '0.00%', 'center')
        d.comment = Comment(why, 'shohiyo')
        put(f'E{r}', f'=IF(C{r}<=0,0,MAX(1000,ROUNDDOWN(ROUNDDOWN(C{r},-3)*D{r},-2)))', 11, False, None, YEN, 'right')
        r += 1
    put(f'B{r}', '司法書士報酬・実費', 11)
    put(f'E{r}', s.get('scrivener_fee', JUDICIAL_SCRIVENER_FEE), 11, False, BLUE, YEN, 'right').comment = \
        Comment('司法書士報酬と登記事項証明書等の実費（税込の目安）', 'shohiyo')
    r += 1
    assert r == reg_total
    put(f'C{r}', '登記費用 計', 11, True, align='right').comment = Comment('上の諸経費には10万円単位で繰り上げて計上', 'shohiyo')
    put(f'E{r}', f'=SUM(E{reg_head + 1}:E{r - 1})', 12, True, None, YEN, 'right')
    for col in 'BCDE':
        ws[f'{col}{r}'].border = Border(top=thin)


    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f'A1:H{print_end}'
    wb.calculation.fullCalcOnLoad = True  # Excel で開いたら必ず再計算
    wb.save(out_path)


def inject_cached_values(path):
    """数式の計算結果を xlsx に書き込む（openpyxl は計算結果を保存しないため、
    スマホのプレビュー等で空欄になるのを防ぐ）。Excel で開けば通常どおり再計算される。
    計算には formulas パッケージを使う。無ければ何もしない。"""
    try:
        import formulas
    except ImportError:
        return False
    warnings.filterwarnings('ignore')
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    values = {}
    for key, v in sol.items():
        ref = key.split('!')[-1]
        if ':' in ref or not hasattr(v, 'value'):
            continue
        val = v.value[0][0]
        try:
            values[ref] = float(val)
        except (TypeError, ValueError):
            if isinstance(val, str) and not val.startswith('#'):
                values[ref] = val
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    sheet = 'xl/worksheets/sheet1.xml'
    xml = files[sheet].decode('utf-8')

    def add_value(m):
        ref, body = m.group(1), m.group(3)
        if ref not in values:
            return m.group(0)
        num = values[ref]
        if isinstance(num, str):
            attrs = re.sub(r' t="[^"]*"', '', m.group(2))
            return f'<c r="{ref}"{attrs} t="str">{body}<v>{escape(num)}</v></c>'
        if abs(num - round(num)) < 1e-6:  # 浮動小数点の誤差（399999.99999…）を整数に
            num = round(num)
        text = str(int(num)) if num == int(num) else repr(num)
        return f'<c r="{ref}"{m.group(2)}>{body}<v>{text}</v></c>'

    xml = re.sub(r'<c r="([A-Z]+[0-9]+)"([^>]*)>(<f>.*?</f>)(?:<v */>|<v></v>)?</c>', add_value, xml)
    files[sheet] = xml.encode('utf-8')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, data in files.items():
            z.writestr(n, data)
    return True


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    spec = normalize(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8')))
    out = sys.argv[2] if len(sys.argv) > 2 else f"資金計算書_{spec.get('name') or '物件'}_{datetime.date.today():%Y%m%d}.xlsx"
    write_workbook(spec, out)
    inject_cached_values(out)
    print(out)


if __name__ == '__main__':
    main()
